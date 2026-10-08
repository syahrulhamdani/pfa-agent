"""Reading bank and card statements.

File parsing (CSV and Excel), batching rows for the model, cleaning up what the model
returns, spotting transactions that were already saved, and the summary shown to the
user. All of it is plain code so it can be tested without a model: the numbers the
user confirms against are never computed by the model.
"""

import asyncio
import csv
import io
from collections.abc import Iterable
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from pfa.core.exceptions import StatementError
from pfa.models.finance import Direction
from pfa.schemas.extraction import ExtractedStatementRow
from pfa.schemas.statements import (
    NormalizedRows,
    ParsedRow,
    Spend,
    StatementFormat,
    StatementSummary,
    StatementTotals,
)
from pfa.services.pending import parse_amount, parse_date

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 500  # data rows, not counting the header
ROWS_PER_BATCH = 100  # rows sent to the model per call, so each answer stays short

_EXTENSIONS: dict[str, StatementFormat] = {
    ".pdf": "pdf",
    ".csv": "csv",
    ".xlsx": "xlsx",
}
_MIME_TYPES: dict[str, StatementFormat] = {
    "application/pdf": "pdf",
    "text/csv": "csv",
    "text/comma-separated-values": "csv",
    "application/csv": "csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}
SUPPORTED_TEXT = "PDF, CSV or Excel (.xlsx)"


def statement_format(
    filename: str | None, mime_type: str | None
) -> StatementFormat | None:
    """Tells which supported statement format a file is, or None if it is not one.

    The extension decides first. The old binary `.xls` format is not supported even
    when its MIME type looks like a spreadsheet.
    """
    extension = Path(filename or "").suffix.lower()
    if extension in _EXTENSIONS:
        return _EXTENSIONS[extension]
    if extension == ".xls":
        return None
    return _MIME_TYPES.get((mime_type or "").lower().split(";")[0].strip())


# --- reading the file into rows of text ---


def parse_csv(content: bytes) -> list[list[str]]:
    """Reads CSV bytes into rows, guessing the delimiter and the text encoding."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = [
            [cell.strip() for cell in row]
            for row in csv.reader(io.StringIO(text, newline=""), dialect)
        ]
    except csv.Error as exc:
        raise StatementError("I couldn't read this CSV file.") from exc
    return [row for row in rows if any(row)]


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time() == datetime.min.time()
            else value.isoformat(" ")
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)).normalize(), "f")
    return str(value).strip()


def parse_xlsx(content: bytes) -> list[list[str]]:
    """Reads the first non-empty sheet of an .xlsx file into rows of text."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # a corrupt or non-Excel file raises many kinds of error
        raise StatementError(
            "I couldn't open this Excel file. Is it a valid .xlsx file?"
        ) from exc
    try:
        for sheet in workbook.worksheets:
            rows = [
                [_cell_text(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(row)]
            if rows:
                return rows
        return []
    finally:
        workbook.close()


async def read_table(content: bytes, fmt: StatementFormat) -> list[list[str]]:
    """Reads a CSV or Excel statement into rows, with size limits.

    Raises:
        StatementError: If the file is too big, has too many rows, or is empty.
    """
    if len(content) > MAX_BYTES:
        raise StatementError(
            f"This file is larger than {MAX_BYTES // (1024 * 1024)} MB. "
            "Please split it into smaller files."
        )
    parser = parse_csv if fmt == "csv" else parse_xlsx
    rows = await asyncio.to_thread(parser, content)
    if len(rows) < 2:
        raise StatementError("I couldn't find any transaction rows in this file.")
    if len(rows) - 1 > MAX_ROWS:
        raise StatementError(
            f"This statement has {len(rows) - 1} rows and the limit is {MAX_ROWS}. "
            "Please split it into smaller files."
        )
    return rows


def table_batches(rows: list[list[str]], size: int = ROWS_PER_BATCH) -> list[str]:
    """Turns rows into tab-separated text batches, repeating the header in each."""

    def clean(cell: str) -> str:
        return cell.replace("\t", " ").replace("\r", " ").replace("\n", " ")

    header, data = rows[0], rows[1:]
    return [
        "\n".join(
            "\t".join(clean(c) for c in row) for row in [header, *data[i : i + size]]
        )
        for i in range(0, len(data), size)
    ]


# --- cleaning up what the model returned ---


def normalize_rows(rows: Iterable[ExtractedStatementRow]) -> NormalizedRows:
    """Keeps complete lines, money in and money out alike, and counts the rest.

    A line missing its merchant, amount or date, or with an amount of zero, is left
    out rather than guessed at. Amounts are saved as positive numbers; the direction
    records whether the money came in or went out.
    """
    kept: list[ParsedRow] = []
    unreadable = 0
    for row in rows:
        merchant = (row.merchant or "").strip()
        amount = parse_amount(row.amount)
        day = parse_date(row.occurred_on)
        if not merchant or amount is None or amount == 0 or day is None:
            unreadable += 1
            continue
        kept.append(
            ParsedRow(
                merchant=merchant,
                amount=abs(amount),
                occurred_on=day,
                direction=Direction(row.direction),
                category=(row.category or "").strip() or None,
                description=(row.description or "").strip() or None,
            )
        )
    return NormalizedRows(rows=kept, skipped_unreadable=unreadable)


# --- duplicates of what is already saved ---


def looks_like_duplicate(candidate: Spend, existing: Spend) -> bool:
    """Decides whether a statement line is a transaction that is already saved.

    The candidate is only ever compared with transactions the user has already
    confirmed on the same card, never with other lines of the same statement, so two
    identical purchases on one day both survive.
    """
    return (
        candidate.direction == existing.direction  # a refund is not its purchase
        and candidate.occurred_on == existing.occurred_on
        and candidate.amount == existing.amount
        and candidate.merchant.strip().casefold()
        == existing.merchant.strip().casefold()
    )


# --- the summary shown before the user confirms ---


def summarize(amounts: Iterable[Decimal]) -> StatementSummary:
    """Computes the summary in code; the model never produces these numbers."""
    values = list(amounts)
    if not values:
        return StatementSummary(
            count=0, total=Decimal("0.00"), highest=None, lowest=None
        )
    return StatementSummary(
        count=len(values),
        total=sum(values, Decimal("0.00")),
        highest=max(values),
        lowest=min(values),
    )


def summarize_by_direction(
    items: Iterable[tuple[Direction, Decimal]],
) -> StatementTotals:
    """Summarises money out and money in separately, in code."""
    out: list[Decimal] = []
    received: list[Decimal] = []
    for direction, amount in items:
        (out if direction is Direction.OUT else received).append(amount)
    return StatementTotals(money_out=summarize(out), money_in=summarize(received))
